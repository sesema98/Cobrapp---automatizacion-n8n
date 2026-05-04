from __future__ import annotations

import base64
import io
import os
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

import cv2
import numpy as np
import pytesseract
from fastapi import FastAPI, File, HTTPException, UploadFile
from pydantic import BaseModel
from openpyxl import Workbook, load_workbook
from PIL import Image


APP_NAME = "CobrApp - API OCR"
PAGOS_DIR = Path(os.getenv("PAGOS_DIR", "/app/pagos"))
EXCEL_PATH = PAGOS_DIR / "pagos.xlsx"

HEADERS = [
    "Fecha",
    "Hora",
    "Nombre",
    "Monto",
    "Tipo",
    "N Operacion",
    "Estado",
    "Registrado por",
    "Texto OCR",
]

app = FastAPI(title=APP_NAME)


class ImagenBase64Request(BaseModel):
    imagen_base64: str


def normalizar_texto(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(caracter for caracter in texto if not unicodedata.combining(caracter))
    return texto.lower()


def asegurar_excel() -> None:
    PAGOS_DIR.mkdir(parents=True, exist_ok=True)

    if EXCEL_PATH.exists():
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pagos"
    sheet.append(HEADERS)
    workbook.save(EXCEL_PATH)


def leer_filas_pago() -> list[dict[str, Any]]:
    asegurar_excel()
    workbook = load_workbook(EXCEL_PATH)
    sheet = workbook.active

    pagos: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        pagos.append(dict(zip(HEADERS, row)))

    return pagos


def existe_operacion(numero_operacion: str | None) -> bool:
    if not numero_operacion:
        return False

    numero_limpio = re.sub(r"\D", "", numero_operacion)
    for pago in leer_filas_pago():
        registrado = re.sub(r"\D", "", str(pago.get("N Operacion") or ""))
        if registrado and registrado == numero_limpio:
            return True

    return False


def registrar_pago(datos: dict[str, Any]) -> dict[str, Any]:
    asegurar_excel()
    estado = "Duplicado" if existe_operacion(datos.get("operacion")) else "Registrado"

    workbook = load_workbook(EXCEL_PATH)
    sheet = workbook.active
    sheet.append(
        [
            datos.get("fecha"),
            datos.get("hora"),
            datos.get("nombre"),
            float(datos["monto"]) if datos.get("monto") else None,
            datos.get("tipo"),
            datos.get("operacion"),
            estado,
            "Bot automatico",
            datos.get("texto_raw"),
        ]
    )
    try:
        workbook.save(EXCEL_PATH)
    except PermissionError as exc:
        return {
            **datos,
            "valido": False,
            "estado": "Excel bloqueado",
            "error": (
                "No se pudo guardar el pago porque pagos.xlsx esta abierto o bloqueado. "
                "Cierra el archivo de Excel y vuelve a enviar la captura."
            ),
        }

    return {**datos, "estado": estado}


def preprocesar_imagen(image_bytes: bytes) -> np.ndarray:
    try:
        image = Image.open(io.BytesIO(image_bytes)).convert("RGB")
    except Exception as exc:
        raise HTTPException(status_code=400, detail="El archivo no es una imagen valida") from exc

    np_image = np.array(image)
    bgr_image = cv2.cvtColor(np_image, cv2.COLOR_RGB2BGR)
    gray = cv2.cvtColor(bgr_image, cv2.COLOR_BGR2GRAY)
    gray = cv2.resize(gray, None, fx=1.7, fy=1.7, interpolation=cv2.INTER_CUBIC)
    gray = cv2.GaussianBlur(gray, (3, 3), 0)
    _, thresh = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    return thresh


def extraer_texto_ocr(image_bytes: bytes) -> str:
    imagen_procesada = preprocesar_imagen(image_bytes)

    try:
        return pytesseract.image_to_string(imagen_procesada, lang="spa+eng")
    except pytesseract.TesseractError:
        return pytesseract.image_to_string(imagen_procesada)


def buscar_monto(texto: str) -> str | None:
    patrones = [
        r"(?:s\/\.?|soles?|pen)\s*([0-9]+(?:[.,][0-9]{1,2})?)",
        r"\bs\s*[\/|iIl1]?\s*([0-9]{1,4}(?:[.,][0-9]{1,2})?)\b",
        r"([0-9]+(?:[.,][0-9]{1,2})?)\s*(?:s\/\.?|soles?)",
        r"(?:monto|importe|total|pago)\D{0,12}([0-9]+(?:[.,][0-9]{1,2})?)",
    ]

    for patron in patrones:
        match = re.search(patron, texto, flags=re.IGNORECASE)
        if match:
            return match.group(1).replace(",", ".")

    return None


def buscar_operacion(texto: str) -> str | None:
    patrones = [
        r"(?:operaci[oó]n|op\.?|codigo|c[oó]digo|transacci[oó]n|movimiento)\D{0,18}(\d{5,})",
        r"\b(\d{8,14})\b",
    ]

    for patron in patrones:
        match = re.search(patron, texto, flags=re.IGNORECASE)
        if match:
            return match.group(1)

    return None


def buscar_fecha(texto: str) -> str:
    match = re.search(r"\b(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\b", texto)
    if not match:
        return date.today().strftime("%d/%m/%Y")

    partes = re.split(r"[/-]", match.group(1))
    dia, mes, anio = partes
    if len(anio) == 2:
        anio = f"20{anio}"
    return f"{int(dia):02d}/{int(mes):02d}/{anio}"


def buscar_hora(texto: str) -> str:
    match = re.search(r"\b(\d{1,2}:\d{2})(?:\s*(?:a\.?\s*m\.?|p\.?\s*m\.?|am|pm))?\b", texto, flags=re.IGNORECASE)
    if match:
        return match.group(1)
    return datetime.now().strftime("%H:%M")


def buscar_tipo(texto_normalizado: str) -> str:
    if "yape" in texto_normalizado:
        return "Yape"
    if "plin" in texto_normalizado:
        return "Plin"
    return "Desconocido"


def buscar_nombre(texto: str) -> str | None:
    lineas = [linea.strip() for linea in texto.splitlines() if linea.strip()]
    token_nombre = r"(?:[A-ZÁÉÍÓÚÑ][A-Za-zÁÉÍÓÚÑáéíóúñ]+|[A-ZÁÉÍÓÚÑ]\.)"
    patron_nombre = re.compile(rf"\b({token_nombre}(?:\s+{token_nombre}){{1,4}})\b")
    palabras_descartadas = re.compile(
        r"(yape|plin|operaci[oó]n|codigo|c[oó]digo|seguridad|datos|transacci[oó]n|"
        r"monto|soles|s\/|fecha|hora|celular|destino|nro\.?)",
        re.IGNORECASE,
    )

    for linea in lineas:
        if palabras_descartadas.search(linea) or re.search(r"\d", linea):
            continue
        match = patron_nombre.search(linea)
        if match:
            return re.sub(r"\s+", " ", match.group(1)).strip()

    return None


def extraer_datos(texto: str) -> dict[str, Any]:
    texto_normalizado = normalizar_texto(texto)
    monto = buscar_monto(texto)
    operacion = buscar_operacion(texto)
    tipo = buscar_tipo(texto_normalizado)
    nombre = buscar_nombre(texto)
    valido = bool(monto and operacion and tipo != "Desconocido")

    return {
        "valido": valido,
        "fecha": buscar_fecha(texto),
        "hora": buscar_hora(texto),
        "nombre": nombre or "No identificado",
        "monto": monto,
        "tipo": tipo,
        "operacion": operacion,
        "texto_raw": texto.strip(),
    }


async def procesar_bytes(image_bytes: bytes, registrar: bool = True) -> dict[str, Any]:
    texto = extraer_texto_ocr(image_bytes)
    datos = extraer_datos(texto)

    if registrar and datos["valido"]:
        datos = registrar_pago(datos)
    elif registrar:
        datos["estado"] = "No registrado"

    return datos


@app.get("/")
def home() -> dict[str, str]:
    return {
        "app": APP_NAME,
        "estado": "ok",
        "docs": "/docs",
    }


@app.post("/procesar-imagen")
async def procesar_imagen(file: UploadFile = File(...)) -> dict[str, Any]:
    contenido = await file.read()
    if not contenido:
        raise HTTPException(status_code=400, detail="La imagen esta vacia")

    return await procesar_bytes(contenido)


@app.post("/procesar-imagen-base64")
async def procesar_imagen_base64(payload: ImagenBase64Request) -> dict[str, Any]:
    try:
        image_bytes = base64.b64decode(payload.imagen_base64.split(",")[-1], validate=False)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="La imagen_base64 no es valida") from exc

    return await procesar_bytes(image_bytes)


@app.get("/pagos")
def listar_pagos() -> dict[str, Any]:
    pagos = leer_filas_pago()
    return {
        "total": len(pagos),
        "pagos": pagos,
    }


@app.get("/reporte")
def reporte() -> dict[str, Any]:
    hoy = date.today().strftime("%d/%m/%Y")
    pagos_hoy = [
        pago
        for pago in leer_filas_pago()
        if pago.get("Fecha") == hoy and pago.get("Estado") == "Registrado"
    ]
    total = sum(float(pago.get("Monto") or 0) for pago in pagos_hoy)

    return {
        "fecha": hoy,
        "cantidad_pagos": len(pagos_hoy),
        "total_recaudado": round(total, 2),
    }
